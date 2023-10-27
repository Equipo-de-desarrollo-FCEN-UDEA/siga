import json
import os
import re
from datetime import datetime
from uuid import uuid1
from io import BytesIO
from tempfile import NamedTemporaryFile
import tempfile

from docx import Document
from zipfile import ZipFile, ZIP_DEFLATED

from app.domain.models import User, EconomicSupport
from app.domain.schemas import UserResponse
from app.domain.schemas import EconomicSupportResponse
from ..templates import templates_dir
from app.core.config import get_app_settings
from app.core.logging import get_logging
from app.core.celery_worker import celery_app
from app.services import aws

settings = get_app_settings()
log = get_logging(__name__)

def fill_economic_support_form(user: User, economic_support:EconomicSupportResponse):
    ...

#     economic_support_dict: dict = economic_support.dict()
#     log.debug("ESTE ES EL DICCIONARIO...")
#     log.debug(economic_support_dict)

#     user = UserResponse.from_orm(user).dict(exclude_unset=True)

#     #Rellenar datos de la solicitud.
#     today = economic_support_dict['created_at'].strftime("%d-%b-%Y")
#     data_economic_support = {
#         0:{'today_date':today},
#         1:{
#             'type':economic_support_dict['economic_support']['application_data']['application_type'],
#             'applicant_name':' '.join([user['names'], user['last_names']]).title(),
#             'identification':user['identification_number'],
#             'project':economic_support_dict['economic_support']['application_data']['project'],
#             'objective':economic_support_dict['economic_support']['application_data']['goal'],
#         },
#         2:{
#             'for_vinculado': economic_support_dict['economic_support']['personal_data']['application_for'] if (economic_support_dict['economic_support']['personal_data']['application_for']=='Vinculado UdeA') else (''), #ó Personal externo/Estudiante/Contratista.
#             'for_external':economic_support_dict['economic_support']['personal_data']['application_for'] if (economic_support_dict['economic_support']['personal_data']['application_for']!='Vinculado UdeA') else (''),
#             'name':economic_support_dict['economic_support']['personal_data']['name'],
#             'identification_number':economic_support_dict['economic_support']['personal_data']['identification_number'],
#             'address':economic_support_dict['economic_support']['personal_data']['address'],
#             'city':economic_support_dict['economic_support']['personal_data']['city'],
#             'telephone':economic_support_dict['economic_support']['personal_data']['landline'],
#             'email':economic_support_dict['economic_support']['personal_data']['email'],
#             'phone':economic_support_dict['economic_support']['personal_data']['phone'],
#             'institution':economic_support_dict['economic_support']['personal_data']['institution'],
#             'academic_unit':economic_support_dict['economic_support']['personal_data']['academic_unit'],
#         },
#         3:{
#             'birth_date':'',
#             'departure_date':economic_support_dict['economic_support']['tickets']['departure_date'].strftime("%d-%b-%Y"),
#             'departure_place':economic_support_dict['economic_support']['tickets']['departure_place'],
#             'birth_place':economic_support_dict['economic_support']['tickets']['place_birth'],
#             'arrival_date':economic_support_dict['economic_support']['tickets']['arrival_date'].strftime("%d-%b-%Y"),
#             'arrival_place':economic_support_dict['economic_support']['tickets']['arrival_place']
#         },
#         4:{
#             'payment_name':economic_support_dict['economic_support']['payment']['name'],
#             'identification':economic_support_dict['economic_support']['payment']['id'],
#             'bank':economic_support_dict['economic_support']['payment']['bank'],
#             'value':economic_support_dict['economic_support']['payment']['value'],
#             'account_number':economic_support_dict['economic_support']['payment']['account_number'],
#             'account_type':economic_support_dict['economic_support']['payment']['account_type'],#corriente
#             'start_date':economic_support_dict['economic_support']['payment']['start_date'].strftime("%d-%b-%Y"),
#             'end_date':economic_support_dict['economic_support']['payment']['end_date'].strftime("%d-%b-%Y")
#         }
#     }

#     path = f'user_{user["id"]}/{uuid1()}'+'formato_apoyo_economico.docx'

#     generate_support_format_to_aws.apply_async(args=(data_economic_support, path))
#     log.debug(path)
#     return path

# @celery_app.task
# def generate_support_format_to_aws(data_application: dict, path: str):
#     cells = open(templates_dir + '/cells_support.json')
#     cells = json.load(cells)

#     document = Document(templates_dir + '/en-fo-001.v02-solicitud-avances-viaticos-pasajes-apoyos-economicos.docx')

#     for table in range(5):
#         keys, values = cells[str(table)].keys(), cells[str(table)].values()
#         for key, cell in zip(keys, values):
#             document.tables[table].cell(0,cell).text = data_application[str(table)][key]

#     with NamedTemporaryFile() as tmp:
#         document.save(tmp.name)
#         file = BytesIO(tmp.read())

#     aws.s3.push_data_to_s3_bucket(settings.aws_bucket_name, file,
#                                   file_name=path, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

# @celery_app.task
def create_zip_documents(user: User, economic_support:EconomicSupportResponse):
    ...
#     regex = r"user_[0-9]+/([\w.-]+)"
#     economic_support_dict: dict = economic_support.dict()
#     files = economic_support_dict['economic_support']['documents']
#     zip_name = "{}_{}-{}_{}.zip".format(economic_support_dict['application_sub_type']['name'],
#                                   economic_support_dict['id'],
#                                   str.title(economic_support_dict['user']['names']+economic_support_dict['user']['last_names']),
#                                   economic_support_dict['user']['identification_number']
#                                   )
#     file_paths = [file['path'] for file in files]
#     #Abrir una carpeta temporal.
#     temp_dir = tempfile.TemporaryDirectory()
#     # log.debug("Está vacía la carpeta?")
#     # log.debug(os.listdir(temp_dir.name))
#     destination = temp_dir.name
#     binary_files = [os.path.join(destination, re.findall(regex, file)[0]) for file in file_paths]
#     for file in file_paths:
#         aws.s3.s3.meta.client.download_file(settings.aws_bucket_name, file, os.path.join(destination, re.findall(regex, file)[0]))

#     # log.debug("Está llena la carpeta?")
#     # log.debug(os.listdir(temp_dir.name))

#     zip_buffer = BytesIO()
#     with ZipFile(zip_buffer,'w', ZIP_DEFLATED) as zip:
#         # writing each file one by one
#         for file in binary_files:
#             zip.write(file)

#     return zip_buffer
